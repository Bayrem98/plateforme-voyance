"""
Script d'import des fichiers Word (.docx) vers la plateforme.
Détection automatique du thème par mots-clés dans le nom du fichier.
Usage : python import_word.py
"""
import os
import re
import unicodedata
from docx import Document
from app import app
from models import db, Theme, Formation


DOSSIER_WORD = "docs_word"

# 🗺️ Règles de détection : si le nom du fichier contient un mot-clé,
# alors on l'associe au thème correspondant.
REGLES_THEME = [
    (["numerolog"], "Numérologie"),
    (["astrolog", "zodiaque", "signe"], "Astrologie"),
    (["tarot", "marseill"], "Tarot de Marseille"),
    (["taromanc"], "Taromancie"),
     # 🆕 Détection des arcanes du Tarot
    (["bateleur", "arcane", "mat", "papesse", "imperatrice", "empereur",
      "pape", "amoureux", "chariot", "force", "hermite", "roue",
      "justice", "pendu", "mort", "temperance", "diable", "tour",
      "etoile", "lune", "soleil", "jugement", "monde"], "Tarot de Marseille"),
]


def normaliser(txt):
    """Enlève les accents et met en minuscules pour comparaison."""
    txt = unicodedata.normalize('NFD', txt)
    txt = ''.join(c for c in txt if unicodedata.category(c) != 'Mn')
    return txt.lower()


def detecter_theme(nom_fichier):
    """Retourne le nom du thème ou None."""
    nom_norm = normaliser(nom_fichier)
    for mots_cles, theme in REGLES_THEME:
        for mot in mots_cles:
            if mot in nom_norm:
                return theme
    return None


def docx_to_html(path):
    """Convertit un .docx en HTML propre : titres, paragraphes, tableaux."""
    doc = Document(path)
    html = []
    
    # ⚠️ On itère sur les paragraphes ET les tableaux DANS L'ORDRE
    # en parcourant le body XML
    from docx.oxml.ns import qn
    
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split('}')[-1]
        
        if tag == 'p':
            # Trouver le paragraphe correspondant
            for p in doc.paragraphs:
                if p._element is child:
                    style = p.style.name if p.style else ''
                    texte = p.text.strip()
                    if not texte:
                        continue
                    
                    # Échapper le HTML
                    texte_esc = (texte.replace('&', '&amp;')
                                       .replace('<', '&lt;')
                                       .replace('>', '&gt;'))
                    
                    # Détecter les titres
                    style_norm = normaliser(style)
                    if 'heading 1' in style_norm or 'titre 1' in style_norm:
                        html.append(f"<h2>{texte_esc}</h2>")
                    elif 'heading 2' in style_norm or 'titre 2' in style_norm:
                        html.append(f"<h3>{texte_esc}</h3>")
                    elif 'heading 3' in style_norm or 'titre 3' in style_norm:
                        html.append(f"<h4>{texte_esc}</h4>")
                    elif 'title' in style_norm:
                        html.append(f"<h1>{texte_esc}</h1>")
                    else:
                        # Texte en gras ?
                        runs_visibles = [r for r in p.runs if r.text.strip()]
                        if runs_visibles and all(r.bold for r in runs_visibles):
                            html.append(f"<p><strong>{texte_esc}</strong></p>")
                        else:
                            # Détecter les listes à puces (commence par • ou -)
                            if texte.lstrip().startswith(('•', '-', '·')):
                                contenu = texte.lstrip('•-· ').strip()
                                html.append(f"<li>{contenu}</li>")
                            else:
                                html.append(f"<p>{texte_esc}</p>")
                    break
        
        elif tag == 'tbl':
            for table in doc.tables:
                if table._element is child:
                    html.append('<table class="table table-bordered">')
                    for i, row in enumerate(table.rows):
                        html.append('<tr>')
                        for cell in row.cells:
                            contenu = (cell.text.strip()
                                       .replace('&', '&amp;')
                                       .replace('<', '&lt;')
                                       .replace('>', '&gt;')
                                       .replace('\n', '<br>'))
                            balise = 'th' if i == 0 else 'td'
                            html.append(f'<{balise}>{contenu}</{balise}>')
                        html.append('</tr>')
                    html.append('</table>')
                    break
    
    # Fermer les <li> dans des <ul>
    resultat = "\n".join(html)
    resultat = re.sub(r'(<li>.*?</li>\n?)+', 
                      lambda m: '<ul>' + m.group(0) + '</ul>', 
                      resultat, flags=re.DOTALL)
    return resultat


def importer():
    with app.app_context():
        if not os.path.exists(DOSSIER_WORD):
            os.makedirs(DOSSIER_WORD)
            print(f"📁 Dossier '{DOSSIER_WORD}' créé.")
            print(f"👉 Place tes fichiers .docx dedans puis relance le script.")
            return
        
        fichiers = [f for f in os.listdir(DOSSIER_WORD) if f.lower().endswith('.docx')]
        if not fichiers:
            print(f"⚠️  Aucun fichier .docx dans '{DOSSIER_WORD}'.")
            return
        
        print(f"📂 {len(fichiers)} fichier(s) trouvé(s)\n")
        
        for nom_fichier in fichiers:
            chemin = os.path.join(DOSSIER_WORD, nom_fichier)
            theme_nom = detecter_theme(nom_fichier)
            
            if not theme_nom:
                print(f"⚠️  '{nom_fichier}' : aucun thème détecté → ignoré.")
                continue
            
            theme = Theme.query.filter_by(nom=theme_nom).first()
            if not theme:
                print(f"❌ Thème '{theme_nom}' introuvable en base → ignoré.")
                continue
            
            # Titre = nom du fichier sans extension
            titre = os.path.splitext(nom_fichier)[0]
            
            # Vérifier si déjà importé
            existante = Formation.query.filter_by(
                theme_id=theme.id, titre=titre
            ).first()
            
            contenu_html = docx_to_html(chemin)
            
            if existante:
                existante.contenu = contenu_html
                print(f"♻️  Mis à jour : {titre}  [{theme_nom}]")
            else:
                ordre = Formation.query.filter_by(theme_id=theme.id).count() + 1
                f = Formation(
                    titre=titre,
                    contenu=contenu_html,
                    theme_id=theme.id,
                    ordre=ordre
                )
                db.session.add(f)
                print(f"➕ Ajouté : {titre}  [{theme_nom}]")
        
        db.session.commit()
        print("\n✅ Import terminé !")


if __name__ == "__main__":
    importer()